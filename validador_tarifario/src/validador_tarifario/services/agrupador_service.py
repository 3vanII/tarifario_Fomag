from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from pathlib import Path
import re
from typing import Callable
import unicodedata

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

from validador_tarifario.utils.paths import (
    DEFAULT_SCAN_ROWS,
    ERROR_SHEET_NAME,
    SERVICIOS_REPS_FILE,
    SOAT_2025_FILE,
    SOAT_UVB_2026_FILE,
    SUPPORTED_EXTENSIONS,
    TABLA_REFERENCIA_CUPS_FILE,
)

LogFn = Callable[[str], None]


class AgrupadorError(Exception):
    pass


@dataclass(slots=True)
class AgrupadorConfig:
    input_file: Path
    output_dir: Path
    sheet_name: str = "AGRUPADOR"
    scan_rows: int = DEFAULT_SCAN_ROWS


@dataclass(slots=True)
class AgrupadorResult:
    output_file: Path
    total_sheets: int
    processed_sheets: int
    total_rows: int
    error_rows: int
    total_input_rows: int


@dataclass(slots=True)
class RepsReferenceData:
    site_names: dict[str, str]
    site_services: dict[str, set[str]]
    service_names_by_pair: dict[tuple[str, str], str]
    service_names: dict[str, str]


@dataclass(slots=True)
class TariffReferenceData:
    soat_2025: dict[str, float]
    soat_uvb_2026: dict[str, float]


class ExcelAgrupadorService:
    def __init__(self, logger: LogFn | None = None) -> None:
        self.logger = logger or (lambda msg: None)
        self._reference_codes_cache: set[str] | None = None
        self._reps_reference_cache: RepsReferenceData | None = None
        self._tariff_reference_cache: TariffReferenceData | None = None

    def log(self, message: str) -> None:
        self.logger(message)

    def process_file(self, config: AgrupadorConfig) -> AgrupadorResult:
        input_file = config.input_file.expanduser().resolve()
        output_dir = config.output_dir.expanduser().resolve()

        self._validate_input(input_file)
        output_dir.mkdir(parents=True, exist_ok=True)

        self.log(f"Leyendo libro: {input_file.name}")
        xls = pd.ExcelFile(input_file)
        all_sheets = list(xls.sheet_names)
        target_sheets = [s for s in all_sheets if s.strip().lower() != config.sheet_name.strip().lower()]

        if not target_sheets:
            raise AgrupadorError("No hay hojas disponibles para agrupar.")

        dataframes: list[pd.DataFrame] = []

        for sheet_name in target_sheets:
            self.log(f"Procesando hoja: {sheet_name}")
            try:
                df = self._read_normalized_sheet(
                    file_path=input_file,
                    sheet_name=sheet_name,
                    scan_rows=config.scan_rows,
                )
                if df.empty:
                    self.log(f"  - Hoja omitida por no contener datos útiles: {sheet_name}")
                    continue

                df.insert(0, "Pestaña origen", sheet_name)
                dataframes.append(df)
                self.log(f"  - Filas agregadas: {len(df):,}")
            except Exception as exc:  # noqa: BLE001
                self.log(f"  - Error en hoja '{sheet_name}': {exc}")

        if not dataframes:
            raise AgrupadorError("No se encontraron datos útiles en ninguna hoja del archivo.")

        grouped_df = pd.concat(dataframes, ignore_index=True, sort=False)
        grouped_df = self._add_fomag_analytics(grouped_df)
        valid_df, error_df = self._validate_grouped_data(grouped_df)

        output_file = output_dir / input_file.name
        self._write_grouped_sheet(
            input_file=input_file,
            output_file=output_file,
            grouped_df=valid_df,
            error_df=error_df,
            grouped_sheet_name=config.sheet_name,
        )

        self.log(f"Archivo generado correctamente: {output_file}")
        self.log(
            f"Filas válidas en '{config.sheet_name}': {len(valid_df):,} • "
            f"Filas con error en '{ERROR_SHEET_NAME}': {len(error_df):,}"
        )
        return AgrupadorResult(
            output_file=output_file,
            total_sheets=len(all_sheets),
            processed_sheets=len(dataframes),
            total_rows=len(valid_df),
            error_rows=len(error_df),
            total_input_rows=len(grouped_df),
        )

    def _validate_input(self, input_file: Path) -> None:
        if not input_file.exists():
            raise AgrupadorError(f"El archivo no existe: {input_file}")
        if not input_file.is_file():
            raise AgrupadorError("La ruta seleccionada no es un archivo.")
        if input_file.suffix.lower() not in SUPPORTED_EXTENSIONS:
            raise AgrupadorError("Solo se permiten archivos .xlsx o .xlsm")

    @staticmethod
    def _is_empty(value) -> bool:
        if pd.isna(value):
            return True
        if isinstance(value, str) and value.strip() == "":
            return True
        return False

    @staticmethod
    def _clean_text(value) -> str:
        if value is None or pd.isna(value):
            return ""
        return str(value).strip()

    @staticmethod
    def _normalize_label(value) -> str:
        text = ExcelAgrupadorService._clean_text(value)
        if not text:
            return ""
        text = unicodedata.normalize("NFKD", text)
        text = "".join(char for char in text if not unicodedata.combining(char))
        text = re.sub(r"\s+", " ", text)
        return text.strip().upper()

    @staticmethod
    def _normalize_code(value) -> str:
        if value is None or pd.isna(value):
            return ""

        if isinstance(value, int):
            return str(value)

        if isinstance(value, float):
            if value.is_integer():
                return str(int(value))
            return str(value).strip()

        text = str(value).strip()
        if not text:
            return ""

        if re.fullmatch(r"\d+\.0+", text):
            return text.split(".", maxsplit=1)[0]

        return text

    @classmethod
    def _normalize_lookup_code(cls, value) -> str:
        normalized_code = cls._normalize_code(value)
        if re.fullmatch(r"\d+", normalized_code):
            stripped_code = normalized_code.lstrip("0")
            return stripped_code or "0"
        return normalized_code

    @staticmethod
    def _normalize_amount(value) -> float | None:
        if value is None or pd.isna(value):
            return None

        if isinstance(value, int):
            return float(value)

        if isinstance(value, float):
            return None if pd.isna(value) else float(value)

        text = str(value).strip()
        if not text:
            return None

        text = text.replace("$", "").replace(" ", "")
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        elif "," in text:
            integer_part, _, decimal_part = text.partition(",")
            if decimal_part.isdigit() and 0 < len(decimal_part) <= 2:
                text = f"{integer_part}.{decimal_part}"
            else:
                text = text.replace(",", "")

        try:
            return float(text)
        except ValueError:
            return None

    def _make_unique_columns(self, columns) -> list[str]:
        unique_names: list[str] = []
        counts: dict[str, int] = {}

        for index, column in enumerate(columns, start=1):
            name = self._clean_text(column)
            if not name:
                name = f"COL_{index}"

            name = re.sub(r"\s+", " ", name).strip()

            if name in counts:
                counts[name] += 1
                final_name = f"{name}_{counts[name]}"
            else:
                counts[name] = 1
                final_name = name

            unique_names.append(final_name)

        return unique_names

    def _detect_header_row(self, file_path: Path, sheet_name: str, scan_rows: int) -> int:
        preview = pd.read_excel(
            file_path,
            sheet_name=sheet_name,
            header=None,
            nrows=scan_rows,
            dtype=object,
        )

        best_row = 0
        best_score = -1

        for index, row in preview.iterrows():
            values = [v for v in row.tolist() if not self._is_empty(v)]
            if not values:
                continue

            non_empty_count = len(values)
            text_count = sum(1 for v in values if isinstance(v, str))
            unique_count = len({str(v).strip().lower() for v in values})

            score = (non_empty_count * 10) + (text_count * 3) + unique_count
            if score > best_score:
                best_score = score
                best_row = int(index)

        return best_row

    def _read_normalized_sheet(self, file_path: Path, sheet_name: str, scan_rows: int) -> pd.DataFrame:
        header_row = self._detect_header_row(file_path, sheet_name, scan_rows)

        df = pd.read_excel(
            file_path,
            sheet_name=sheet_name,
            header=header_row,
            dtype=object,
        )

        df = df.dropna(axis=0, how="all").dropna(axis=1, how="all")

        if df.empty:
            return pd.DataFrame()

        df.columns = self._make_unique_columns(df.columns)
        df = df.dropna(axis=0, how="all")
        return df

    def _load_reference_file(self, file_path: Path) -> pd.DataFrame:
        suffix = file_path.suffix.lower()
        if suffix == ".csv":
            return pd.read_csv(file_path, dtype=object, sep=None, engine="python")
        return pd.read_excel(file_path, dtype=object)

    def _get_reference_codes(self) -> set[str]:
        if self._reference_codes_cache is not None:
            return self._reference_codes_cache

        reference_sources = [
            (SOAT_UVB_2026_FILE, "CODIGO"),
            (SOAT_2025_FILE, "CÓDIGO"),
            (TABLA_REFERENCIA_CUPS_FILE, "Codigo"),
        ]

        reference_codes: set[str] = set()

        for file_path, column_name in reference_sources:
            if not file_path.exists():
                raise AgrupadorError(f"No se encontró la base de referencia: {file_path}")

            self.log(f"Cargando base de referencia: {file_path.name}")
            df = self._load_reference_file(file_path)

            matched_column = self._find_column_name(df, column_name)
            if matched_column is None:
                raise AgrupadorError(
                    f"La columna '{column_name}' no existe en la base de referencia {file_path.name}."
                )

            reference_codes.update(
                normalized_code
                for normalized_code in (self._normalize_code(value) for value in df[matched_column].tolist())
                if normalized_code
            )

        self._reference_codes_cache = reference_codes
        return reference_codes

    def _build_reference_amount_map(
        self,
        file_path: Path,
        code_column_name: str,
        value_column_name: str,
    ) -> dict[str, float]:
        if not file_path.exists():
            raise AgrupadorError(f"No se encontró la base de referencia: {file_path}")

        self.log(f"Cargando base de referencia analítica: {file_path.name}")
        df = self._load_reference_file(file_path)

        code_column = self._find_column_name(df, code_column_name)
        value_column = self._find_column_name(df, value_column_name)

        if code_column is None or value_column is None:
            missing_columns = []
            if code_column is None:
                missing_columns.append(f"'{code_column_name}'")
            if value_column is None:
                missing_columns.append(f"'{value_column_name}'")
            joined_columns = ", ".join(missing_columns)
            raise AgrupadorError(
                f"Faltan columnas en la base {file_path.name}: {joined_columns}."
            )

        amount_map: dict[str, float] = {}
        for _, row in df.iterrows():
            code = self._normalize_lookup_code(row[code_column])
            amount = self._normalize_amount(row[value_column])
            if code and amount is not None:
                amount_map[code] = amount

        return amount_map

    def _get_tariff_reference_data(self) -> TariffReferenceData:
        if self._tariff_reference_cache is not None:
            return self._tariff_reference_cache

        self._tariff_reference_cache = TariffReferenceData(
            soat_2025=self._build_reference_amount_map(
                file_path=SOAT_2025_FILE,
                code_column_name="CÓDIGO",
                value_column_name="SOAT- pleno",
            ),
            soat_uvb_2026=self._build_reference_amount_map(
                file_path=SOAT_UVB_2026_FILE,
                code_column_name="CODIGO",
                value_column_name="SOAT UVB 2026",
            ),
        )
        return self._tariff_reference_cache

    def _set_or_insert_column(
        self,
        dataframe: pd.DataFrame,
        column_name: str,
        values: list[float | None],
        insert_at: int,
    ) -> None:
        existing_column = self._find_column_name(dataframe, column_name)
        if existing_column is not None:
            dataframe[existing_column] = values
            return

        dataframe.insert(insert_at, column_name, values)

    def _add_fomag_analytics(self, grouped_df: pd.DataFrame) -> pd.DataFrame:
        cups_column = self._get_required_column(grouped_df, "CUPS")
        tarifa_fomag_column = self._get_required_column(grouped_df, "TARIFA FOMAG")
        tariff_reference = self._get_tariff_reference_data()

        cups_codes = grouped_df[cups_column].map(self._normalize_lookup_code)
        tarifa_fomag_values = grouped_df[tarifa_fomag_column].map(self._normalize_amount)

        difference_vs_soat_2025: list[float | None] = []
        difference_vs_soat_uvb_2026: list[float | None] = []

        matches_soat_2025 = 0
        matches_soat_uvb_2026 = 0

        for row_index in grouped_df.index:
            cups_code = cups_codes.at[row_index]
            tarifa_fomag = tarifa_fomag_values.at[row_index]
            soat_2025 = tariff_reference.soat_2025.get(cups_code)
            soat_uvb_2026 = tariff_reference.soat_uvb_2026.get(cups_code)

            if tarifa_fomag is not None and soat_2025 is not None:
                difference_vs_soat_2025.append(soat_2025 - tarifa_fomag)
                matches_soat_2025 += 1
            else:
                difference_vs_soat_2025.append(None)

            if tarifa_fomag is not None and soat_uvb_2026 is not None:
                difference_vs_soat_uvb_2026.append(soat_uvb_2026 - tarifa_fomag)
                matches_soat_uvb_2026 += 1
            else:
                difference_vs_soat_uvb_2026.append(None)

        insert_position = grouped_df.columns.get_loc(tarifa_fomag_column) + 1
        self._set_or_insert_column(
            dataframe=grouped_df,
            column_name="DIFERENCIA VS SOAT 2025",
            values=difference_vs_soat_2025,
            insert_at=insert_position,
        )
        self._set_or_insert_column(
            dataframe=grouped_df,
            column_name="DIFERENCIA VS SOAT UVB 2026",
            values=difference_vs_soat_uvb_2026,
            insert_at=insert_position + 1,
        )

        self.log(
            "Analítica TARIFA FOMAG: "
            f"{matches_soat_2025:,} fila(s) con comparación contra SOAT 2025 y "
            f"{matches_soat_uvb_2026:,} fila(s) con comparación contra SOAT UVB 2026."
        )
        return grouped_df

    def _find_column_name(self, df: pd.DataFrame, expected_name: str) -> str | None:
        expected = self._normalize_label(expected_name)
        for column in df.columns:
            if self._normalize_label(column) == expected:
                return column
        return None

    def _get_required_column(self, df: pd.DataFrame, expected_name: str) -> str:
        column_name = self._find_column_name(df, expected_name)
        if column_name is None:
            raise AgrupadorError(f"No se encontró la columna '{expected_name}' en el archivo a procesar.")
        return column_name

    def _get_reps_reference_data(self) -> RepsReferenceData:
        if self._reps_reference_cache is not None:
            return self._reps_reference_cache

        if not SERVICIOS_REPS_FILE.exists():
            raise AgrupadorError(f"No se encontró la base de referencia: {SERVICIOS_REPS_FILE}")

        self.log(f"Cargando base de referencia: {SERVICIOS_REPS_FILE.name}")
        df = self._load_reference_file(SERVICIOS_REPS_FILE)

        site_code_column = self._find_column_name(df, "codigo de sede")
        service_code_column = self._find_column_name(df, "serv_codigo")
        site_name_column = self._find_column_name(df, "sede_nombre")
        service_name_column = self._find_column_name(df, "serv_nombre")

        missing_columns = [
            column_name
            for column_name, matched_column in (
                ("codigo de sede", site_code_column),
                ("serv_codigo", service_code_column),
                ("sede_nombre", site_name_column),
                ("serv_nombre", service_name_column),
            )
            if matched_column is None
        ]
        if missing_columns:
            joined_columns = ", ".join(f"'{name}'" for name in missing_columns)
            raise AgrupadorError(
                f"Faltan columnas en la base {SERVICIOS_REPS_FILE.name}: {joined_columns}."
            )

        site_names: dict[str, str] = {}
        site_services: dict[str, set[str]] = {}
        service_names_by_pair: dict[tuple[str, str], str] = {}
        service_names: dict[str, str] = {}

        for _, row in df.iterrows():
            site_code = self._normalize_code(row[site_code_column])
            service_code = self._normalize_code(row[service_code_column])
            site_name = self._clean_text(row[site_name_column])
            service_name = self._clean_text(row[service_name_column])

            if not site_code:
                continue

            if site_name and site_code not in site_names:
                site_names[site_code] = site_name

            site_services.setdefault(site_code, set())

            if not service_code:
                continue

            site_services[site_code].add(service_code)
            if service_name:
                service_names_by_pair[(site_code, service_code)] = service_name
                service_names.setdefault(service_code, service_name)

        self._reps_reference_cache = RepsReferenceData(
            site_names=site_names,
            site_services=site_services,
            service_names_by_pair=service_names_by_pair,
            service_names=service_names,
        )
        return self._reps_reference_cache

    def _validate_grouped_data(self, grouped_df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
        cups_column = self._get_required_column(grouped_df, "CUPS")
        site_code_column = self._get_required_column(grouped_df, "CÓDIGO HABILITACIÓN (12 DÍGITOS)")
        service_code_column = self._get_required_column(grouped_df, "COD SERVICIO")

        reference_codes = self._get_reference_codes()
        reps_reference = self._get_reps_reference_data()

        error_reasons = pd.Series("", index=grouped_df.index, dtype=object)

        normalized_cups_codes = grouped_df[cups_column].map(self._normalize_code)
        invalid_cups_mask = ~normalized_cups_codes.isin(reference_codes)
        error_reasons.loc[invalid_cups_mask] = "Codigo CUPS invalido"

        if invalid_cups_mask.any():
            self.log(f"Validación CUPS: {int(invalid_cups_mask.sum()):,} fila(s) marcadas con error.")
        else:
            self.log("Validación CUPS: no se encontraron códigos inválidos.")

        site_codes = grouped_df[site_code_column].map(self._normalize_code)
        service_codes = grouped_df[service_code_column].map(self._normalize_code)

        reps_error_count = 0

        for row_index in grouped_df.index:
            site_code = site_codes.at[row_index]
            service_code = service_codes.at[row_index]
            current_reason = error_reasons.at[row_index]
            reps_reason = ""

            if not site_code:
                reps_reason = "No se informó el código de habilitación de la sede en la columna 'CÓDIGO HABILITACIÓN (12 DÍGITOS)'."
            elif site_code not in reps_reference.site_services:
                reps_reason = (
                    f"El código de habilitación '{site_code}' no existe en la base REPS "
                    f"({SERVICIOS_REPS_FILE.name})."
                )
            elif not service_code:
                site_name = reps_reference.site_names.get(site_code, "Sede sin nombre reportado en REPS")
                reps_reason = (
                    f"La sede '{site_name}' con código de habilitación '{site_code}' existe en REPS, "
                    "pero no se informó el código en la columna 'COD SERVICIO'."
                )
            elif service_code not in reps_reference.site_services[site_code]:
                site_name = reps_reference.site_names.get(site_code, "Sede sin nombre reportado en REPS")
                service_name = reps_reference.service_names_by_pair.get((site_code, service_code)) or reps_reference.service_names.get(
                    service_code,
                    "",
                )
                if service_name:
                    reps_reason = (
                        f"La sede '{site_name}' con código de habilitación '{site_code}' sí existe en REPS, "
                        f"pero no tiene habilitado el servicio '{service_name}' con código '{service_code}'."
                    )
                else:
                    reps_reason = (
                        f"La sede '{site_name}' con código de habilitación '{site_code}' sí existe en REPS, "
                        f"pero no tiene habilitado el código de servicio '{service_code}'. "
                        "Además, ese código no aparece identificado en la base REPS para traer el nombre del servicio."
                    )

            if reps_reason:
                reps_error_count += 1
                error_reasons.at[row_index] = f"{current_reason} | {reps_reason}" if current_reason else reps_reason

        if reps_error_count:
            self.log(f"Validación REPS: {reps_error_count:,} fila(s) marcadas con error.")
        else:
            self.log("Validación REPS: no se encontraron inconsistencias entre sede y servicio.")

        invalid_mask = error_reasons != ""
        error_df = grouped_df.loc[invalid_mask].copy()
        if not error_df.empty:
            error_df["Motivo del error"] = error_reasons.loc[invalid_mask].values
            self.log(f"Total validaciones: {len(error_df):,} fila(s) enviadas a la hoja de errores.")

        valid_df = grouped_df.loc[~invalid_mask].copy()
        return valid_df, error_df

    @staticmethod
    def _clean_excel_value(value):
        if pd.isna(value):
            return None
        return value

    def _apply_difference_styles(self, worksheet) -> None:
        headers = {
            self._normalize_label(cell.value): cell.column
            for cell in worksheet[1]
            if self._clean_text(cell.value)
        }

        for column_name in ("DIFERENCIA VS SOAT 2025", "DIFERENCIA VS SOAT UVB 2026"):
            column_index = headers.get(self._normalize_label(column_name))
            if column_index is None:
                continue

            for row_index in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row=row_index, column=column_index)
                numeric_value = self._normalize_amount(cell.value)
                if numeric_value is None:
                    continue

                updated_font = copy(cell.font)
                updated_font.color = "FFFF0000" if numeric_value < 0 else "FF000000"
                cell.font = updated_font
                cell.number_format = '#,##0.00;#,##0.00'

    def _append_dataframe_to_worksheet(self, worksheet, dataframe: pd.DataFrame) -> None:
        for row in dataframe_to_rows(dataframe, index=False, header=True):
            worksheet.append([self._clean_excel_value(value) for value in row])

        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions

        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                cell_value = "" if cell.value is None else str(cell.value)
                if len(cell_value) > max_length:
                    max_length = len(cell_value)
            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 40)

    def _write_grouped_sheet(
        self,
        input_file: Path,
        output_file: Path,
        grouped_df: pd.DataFrame,
        error_df: pd.DataFrame,
        grouped_sheet_name: str,
    ) -> None:
        keep_vba = output_file.suffix.lower() == ".xlsm"
        workbook = load_workbook(input_file, keep_vba=keep_vba)

        target_name = grouped_sheet_name.strip() or "AGRUPADOR"

        for sheet_name in list(workbook.sheetnames):
            workbook.remove(workbook[sheet_name])

        worksheet = workbook.create_sheet(title=target_name)
        self._append_dataframe_to_worksheet(worksheet, grouped_df)
        self._apply_difference_styles(worksheet)

        if not error_df.empty:
            error_sheet_name = ERROR_SHEET_NAME
            if error_sheet_name == target_name:
                error_sheet_name = f"{ERROR_SHEET_NAME}_1"
            error_worksheet = workbook.create_sheet(title=error_sheet_name)
            self._append_dataframe_to_worksheet(error_worksheet, error_df)
            self._apply_difference_styles(error_worksheet)

        workbook.save(output_file)
